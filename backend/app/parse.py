"""Pembacaan berkas unggahan menjadi potongan <= 5 SKU. Deterministik, tanpa AI, tanpa menyimpan berkas."""
from __future__ import annotations

import csv
import io
from collections import OrderedDict

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.harness.baseline import map_header
from app.harness.common import norm_key

MAX_BYTES = 2 * 1024 * 1024
MAX_DATA_ROWS = 200
CHUNK_SKUS = 5
MAX_SCAN_ROWS = 5000
EXTRA_SHEET_ROWS = 60
CATEGORIES = ("fashion", "electronics", "beauty")


class ParseError(ValueError):
    """Galat yang pesannya aman ditampilkan ke pengguna."""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def _trim(grid: dict[int, list[str]]) -> dict[int, list[str]]:
    last = 0
    for rn, cells in grid.items():
        if any(c.strip() for c in cells):
            last = max(last, rn)
    return {rn: cells for rn, cells in grid.items() if rn <= last}


def read_xlsx(data: bytes) -> dict[str, dict[int, list[str]]]:
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise ParseError("Berkas XLSX tidak dapat dibaca. Pastikan formatnya .xlsx yang valid.") from e
    sheets: dict[str, dict[int, list[str]]] = {}
    for ws in wb.worksheets:
        grid: dict[int, list[str]] = {}
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i > MAX_SCAN_ROWS:
                raise ParseError(f"Sheet '{ws.title}' terlalu panjang (lebih dari {MAX_SCAN_ROWS} baris).")
            grid[i] = [_cell(v) for v in row]
        grid = _trim(grid)
        if grid:
            width = max(len(r) for r in grid.values())
            sheets[ws.title] = {rn: r + [""] * (width - len(r)) for rn, r in grid.items()}
    wb.close()
    return sheets


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def read_csv(data: bytes) -> dict[int, list[str]]:
    text = _decode(data)
    best: tuple[int, str] = (-1, ",")
    for delim in (",", ";", "\t"):
        rows = list(csv.reader(io.StringIO(text), delimiter=delim))[:15]
        score = max((sum(1 for c in r if map_header(c)) for r in rows), default=0)
        if score > best[0]:
            best = (score, delim)
    try:
        rows = list(csv.reader(io.StringIO(text), delimiter=best[1]))
    except csv.Error as e:
        raise ParseError("Berkas CSV tidak dapat dibaca.") from e
    if len(rows) > MAX_SCAN_ROWS:
        raise ParseError(f"Berkas terlalu panjang (lebih dari {MAX_SCAN_ROWS} baris).")
    grid = _trim({i + 1: r for i, r in enumerate(rows)})
    if grid:
        width = max(len(r) for r in grid.values())
        grid = {rn: r + [""] * (width - len(r)) for rn, r in grid.items()}
    return grid


def find_header(grid: dict[int, list[str]]) -> tuple[int, int | None, list[str | None]]:
    """Baris dengan header terkenal terbanyak (aturan sama dengan baseline.predict)."""
    best: tuple[int, int | None, list[str | None]] = (0, None, [])
    for rn, cells in sorted(grid.items()):
        m = [map_header(c) for c in cells]
        n = sum(1 for x in m if x)
        if n > best[0]:
            best = (n, rn, m)
    return best


def grid_text(grid: dict[int, list[str]], rows: list[int]) -> str:
    ncol = max((len(grid[r]) for r in rows), default=0)
    head = "ROW\t" + "\t".join(get_column_letter(i + 1) for i in range(ncol))
    return head + "\n" + "\n".join(f"{r}\t" + "\t".join(grid[r]) for r in rows)


def infer_category(keys: set[str]) -> str | None:
    if keys & {"warranty", "mpn", "specs"}:
        return "electronics"
    if keys & {"ingredients", "volume"}:
        return "beauty"
    if keys & {"color", "size", "material"}:
        return "fashion"
    return None


def parse_file(filename: str, data: bytes, category: str | None = None) -> dict:
    """Kembalikan {chunks, warnings, ...}. Parser aturan dijalankan terpisah (app.rules)."""
    if len(data) > MAX_BYTES:
        raise ParseError("Berkas lebih dari 2 MB. Kecilkan berkas atau pecah menjadi beberapa bagian.")
    if category is not None and category not in CATEGORIES:
        raise ParseError(f"Kategori tidak dikenal: {category!r}. Pilih salah satu dari {', '.join(CATEGORIES)}.")
    name = filename.lower()
    warnings: list[str] = []
    extra_sheets: dict[str, dict[int, list[str]]] = {}
    if name.endswith(".xlsx"):
        fmt = "xlsx"
        sheets = read_xlsx(data)
        if not sheets:
            raise ParseError("Berkas XLSX kosong.")
        scored = [(find_header(g)[0], sn) for sn, g in sheets.items()]
        _, sheet = max(scored, key=lambda t: t[0])
        grid = sheets[sheet]
        extra_sheets = {sn: g for sn, g in sheets.items() if sn != sheet}
    elif name.endswith(".csv"):
        fmt, sheet = "csv", None
        grid = read_csv(data)
    else:
        raise ParseError("Format berkas tidak didukung. Gunakan .xlsx atau .csv.")
    if not grid:
        raise ParseError("Berkas kosong.")

    n_hdr, hdr_row, cmap = find_header(grid)
    if hdr_row is None or n_hdr < 3 or "sku" not in cmap:
        raise ParseError(
            "Header kolom tidak dikenali. Pastikan ada baris judul kolom yang memuat minimal "
            "kolom SKU dan dua kolom lain (mis. Nama Produk, Merek)."
        )
    col: dict[str, int] = {}
    for i, k in enumerate(cmap):
        if k and k not in col:
            col[k] = i
    unknown = [h.strip() for h, k in zip(grid[hdr_row], cmap, strict=False) if h.strip() and not k]
    if unknown:
        warnings.append("Kolom tidak dikenali dan tidak dibaca parser aturan: " + ", ".join(unknown) + ".")

    sku_col = col["sku"]
    by_key: OrderedDict[str, list[int]] = OrderedDict()
    raw_sku: dict[str, str] = {}
    skipped_nosku: list[int] = []
    multi_variant: list[int] = []
    for rn in sorted(grid):
        if rn <= hdr_row:
            continue
        cells = grid[rn]
        sku = cells[sku_col].strip() if sku_col < len(cells) else ""
        if not sku:
            if sum(1 for c in cells if c.strip()) >= 3:
                skipped_nosku.append(rn)
            continue
        if ";" in sku:
            multi_variant.append(rn)
        key = norm_key(sku)
        by_key.setdefault(key, []).append(rn)
        raw_sku.setdefault(key, sku)
    if not by_key:
        raise ParseError("Tidak ada baris data dengan kode SKU.")
    n_rows = sum(len(v) for v in by_key.values())
    if n_rows > MAX_DATA_ROWS:
        raise ParseError(f"Berkas memuat {n_rows} baris data; batas demo adalah {MAX_DATA_ROWS} baris.")
    if skipped_nosku:
        warnings.append(
            f"{len(skipped_nosku)} baris dilewati karena tanpa kode SKU (baris {_fmt_rows(skipped_nosku)})."
        )
    if multi_variant:
        warnings.append(
            f"{len(multi_variant)} baris memuat beberapa SKU dalam satu sel (baris {_fmt_rows(multi_variant)}); "
            "parser aturan melewatinya, hanya AI yang dapat memecahnya."
        )
    dups = {k: len(v) for k, v in by_key.items() if len(v) > 1}
    if dups:
        warnings.append(
            "SKU muncul lebih dari sekali: " + ", ".join(f"{raw_sku[k]} (x{n})" for k, n in dups.items()) + "."
        )
    for sn in extra_sheets:
        warnings.append(
            f"Sheet '{sn}' disertakan sebagai konteks untuk AI; parser aturan hanya membaca sheet '{sheet}'."
        )

    keys = list(by_key)
    last_data_row = max(max(v) for v in by_key.values())
    tail_rows = [r for r in sorted(grid) if r > last_data_row]
    ctx_rows = list(range(1, hdr_row + 1))
    extra_text = ""
    for sn, g in extra_sheets.items():
        rows = sorted(g)[:EXTRA_SHEET_ROWS]
        extra_text += f"\n\nSHEET: {sn}\n" + grid_text(g, rows)

    chunks: list[dict] = []
    for ci, start in enumerate(range(0, len(keys), CHUNK_SKUS)):
        group = keys[start : start + CHUNK_SKUS]
        rows = sorted(r for k in group for r in by_key[k])
        is_last = start + CHUNK_SKUS >= len(keys)
        shown = ctx_rows + rows + (tail_rows if is_last else [])
        sheet_line = f" | SHEET: {sheet}" if sheet else ""
        text = f"FILE: {filename}{sheet_line}\n" + grid_text(grid, shown) + extra_text
        chunks.append(
            {
                "chunk_id": f"c{ci}",
                "index": ci,
                "skus": [raw_sku[k] for k in group],
                "rows": rows,
                "duplicate_counts": {raw_sku[k]: len(by_key[k]) for k in group if len(by_key[k]) > 1},
                "header_row": hdr_row,
                "columns": col,
                "sheet": sheet,
                "text": text,
                "grid": {"header_row": hdr_row, "rows": {str(r): grid[r] for r in shown}},
            }
        )

    resolved = category or infer_category(set(col))
    return {
        "filename": filename,
        "format": fmt,
        "sheet": sheet,
        "category": resolved,
        "category_source": "user" if category else ("inferred" if resolved else None),
        "n_rows": n_rows,
        "n_skus": len(keys),
        "chunks": chunks,
        "warnings": warnings,
    }


def _fmt_rows(rows: list[int], limit: int = 8) -> str:
    s = ", ".join(str(r) for r in rows[:limit])
    return s + (", ..." if len(rows) > limit else "")
